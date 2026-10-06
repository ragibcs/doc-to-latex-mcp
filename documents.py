import json
import os
import re
import zipfile
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

MAX_PAGES = int(os.environ.get('MAX_PAGES', '40'))
Image.MAX_IMAGE_PIXELS = 25_000_000


def escape_tex(value):
    replacements = {'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}', '$': r'\$', '&': r'\&', '#': r'\#', '%': r'\%', '_': r'\_', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(replacements.get(char, char) for char in str(value) if char in '\n\t' or ord(char) >= 32)


def source_file(job):
    sources = list(job.glob('source.*'))
    if len(sources) != 1:
        raise ValueError('Source document not found.')
    return sources[0]


def pdf_document(job):
    path = source_file(job)
    if path.suffix != '.pdf':
        raise ValueError('This operation requires a PDF.')
    doc = fitz.open(path)
    if doc.needs_pass:
        doc.close()
        raise ValueError('Password-protected PDFs are not supported. Upload an unlocked copy.')
    if not 1 <= len(doc) <= MAX_PAGES:
        doc.close()
        raise ValueError(f'PDF must contain 1–{MAX_PAGES} pages.')
    for page in doc:
        if min(page.rect.width, page.rect.height) < 20 or max(page.rect.width, page.rect.height) > 2000:
            doc.close()
            raise ValueError('Unsupported PDF page dimensions.')
    return doc


def docx_document(job):
    path = source_file(job)
    with zipfile.ZipFile(path) as archive:
        info = archive.infolist()
        if len(info) > 5000 or sum(item.file_size for item in info) > 100 * 1024 * 1024:
            raise ValueError('DOCX expanded contents exceed safety limits.')
        if 'word/document.xml' not in archive.namelist():
            raise ValueError('Not a valid DOCX document.')
    return Document(path)


def inspect(job):
    if source_file(job).suffix == '.docx':
        doc = docx_document(job)
        return {'format': 'docx', 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'images': len(doc.inline_shapes), 'pages': None, 'warnings': ['DOCX pagination is not available without a Word rendering engine. Flowing reconstruction preserves basic headings, text, and tables, not original pagination.']}
    with pdf_document(job) as doc:
        pages = []
        fonts = set()
        for page in doc:
            fonts.update(font[3] for font in page.get_fonts())
            pages.append({'page': page.number + 1, 'width': page.rect.width, 'height': page.rect.height, 'words': len(page.get_text('words')), 'images': len(page.get_images())})
        scanned = [page['page'] for page in pages if page['words'] < 10]
        return {'format': 'pdf', 'pages': len(doc), 'page_details': pages, 'fonts': sorted(fonts), 'warnings': [f'Pages {scanned} may require OCR. OCR is not included; facsimile mode preserves appearance.'] if scanned else []}


def analyze(job, page_number=None):
    if source_file(job).suffix == '.docx':
        doc = docx_document(job)
        result = {'format': 'docx', 'sections': [{'width': s.page_width.pt, 'height': s.page_height.pt, 'left_margin': s.left_margin.pt, 'right_margin': s.right_margin.pt} for s in doc.sections], 'paragraphs': [{'text': p.text, 'style': p.style.name} for p in doc.paragraphs]}
    else:
        pages = []
        with pdf_document(job) as doc:
            if page_number is not None and not 1 <= page_number <= len(doc):
                raise ValueError('Page is out of range.')
            for page in doc:
                if page_number is not None and page.number + 1 != page_number:
                    continue
                blocks = []
                text_blocks = []
                for block in page.get_text('dict', flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)['blocks']:
                    if block['type'] != 0:
                        continue
                    spans = [span for line in block['lines'] for span in line['spans']]
                    blocks.append({'type': 'text', 'bbox': block['bbox'], 'spans': spans})
                    text_blocks.append(fitz.Rect(block['bbox']))
                left = sum(box.x1 < page.rect.width * .56 for box in text_blocks)
                right = sum(box.x0 > page.rect.width * .44 for box in text_blocks)
                columns = 2 if left >= 3 and right >= 3 else 1
                margins = {'left': min((b.x0 for b in text_blocks), default=0), 'top': min((b.y0 for b in text_blocks), default=0), 'right': page.rect.width - max((b.x1 for b in text_blocks), default=page.rect.width), 'bottom': page.rect.height - max((b.y1 for b in text_blocks), default=page.rect.height)}
                pages.append({'page': page.number + 1, 'width': page.rect.width, 'height': page.rect.height, 'estimated_columns': columns, 'estimated_margins': margins, 'blocks': blocks})
        result = {'format': 'pdf', 'pages': pages, 'note': 'Columns and margins are heuristics. Span positions are extracted from the PDF.'}
    (job / 'layout.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    return result


def assets(job):
    target = job / 'figures'
    target.mkdir(exist_ok=True)
    results = []
    if source_file(job).suffix == '.docx':
        doc = docx_document(job)
        for relationship in doc.part.rels.values():
            if 'image' not in relationship.reltype or relationship.is_external:
                continue
            if len(results) >= 100:
                break
            import io
            try:
                with Image.open(io.BytesIO(relationship.target_part.blob)) as image:
                    image.load()
                    name = f'figures/figure-{len(results) + 1:03}.png'
                    image.convert('RGB').save(job / name)
                    results.append({'file': name})
            except (OSError, ValueError):
                continue
    else:
        seen = {}
        with pdf_document(job) as doc:
            for page in doc:
                for info in page.get_image_info(xrefs=True):
                    xref = info['xref']
                    if not xref or info['width'] * info['height'] > 25_000_000:
                        continue
                    if xref not in seen:
                        if len(seen) >= 100:
                            continue
                        name = f'figures/figure-{len(seen) + 1:03}.png'
                        pix = fitz.Pixmap(doc, xref)
                        if pix.colorspace and pix.colorspace.n > 3:
                            pix = fitz.Pixmap(fitz.csRGB, pix)
                        pix.save(job / name)
                        seen[xref] = name
                    results.append({'file': seen[xref], 'page': page.number + 1, 'bbox': info['bbox']})
    (job / 'assets.json').write_text(json.dumps(results), encoding='utf-8')
    return results


def table_tex(rows):
    width = max((len(row) for row in rows), default=0)
    if not width:
        return ''
    body = [' & '.join(escape_tex(cell or '') for cell in row + [''] * (width - len(row))) + r' \\' for row in rows]
    return '\n'.join([r'\begin{tabular}{' + 'l' * width + '}', r'\hline', *body, r'\hline', r'\end{tabular}'])


def tables(job):
    output = []
    if source_file(job).suffix == '.docx':
        for table in docx_document(job).tables:
            output.append({'rows': [[cell.text for cell in row.cells] for row in table.rows]})
    else:
        with pdf_document(job) as doc:
            for page in doc:
                for table in page.find_tables().tables:
                    output.append({'page': page.number + 1, 'bbox': table.bbox, 'rows': table.extract()})
    target = job / 'tables'
    target.mkdir(exist_ok=True)
    for i, table in enumerate(output[:100]):
        name = f'tables/table-{i + 1:03}.tex'
        (job / name).write_text(table_tex(table['rows']), encoding='utf-8')
        table['latex_file'] = name
    (job / 'tables.json').write_text(json.dumps(output), encoding='utf-8')
    return {'tables': output, 'note': 'Detected tables need review, especially merged cells. PDF reconstruction keeps positioned text; exported tabular files are provided separately.'}


def equations(job):
    if source_file(job).suffix == '.docx':
        return {'candidates': [], 'warning': 'DOCX OMML equations are not converted. Review the source document manually.'}
    target = job / 'equations'
    target.mkdir(exist_ok=True)
    output = []
    with pdf_document(job) as doc:
        for page in doc:
            for block in page.get_text('dict', flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)['blocks']:
                if block['type'] != 0:
                    continue
                for line in block['lines']:
                    if len(output) >= 100:
                        break
                    if any(re.search(r'math|symbol|cmsy|cmmi|cmex', span['font'], re.I) for span in line['spans']):
                        name = f'equations/candidate-{len(output) + 1:03}.png'
                        rect = fitz.Rect(line['bbox']) & page.rect
                        if rect.is_empty:
                            continue
                        page.get_pixmap(clip=rect, matrix=fitz.Matrix(1.5, 1.5)).save(job / name)
                        output.append({'page': page.number + 1, 'bbox': list(rect), 'image': name, 'recognized_latex': None})
    (job / 'equations.json').write_text(json.dumps(output), encoding='utf-8')
    return {'candidates': output, 'warning': 'Font-based candidates only, not equation recognition. Ask the LibreChat vision model or a human to transcribe these crops.'}


def docx_body(job):
    doc = docx_document(job)
    output = []
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            content = ''.join((r'\textbf{' + escape_tex(run.text) + '}' if run.bold else r'\textit{' + escape_tex(run.text) + '}' if run.italic else escape_tex(run.text)) for run in block.runs)
            style = block.style.name.lower()
            if style.startswith('heading'):
                depth = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                command = ['section', 'subsection', 'subsubsection'][min(depth, 3) - 1]
                output.append('\\' + command + '*{' + content + '}')
            else:
                output.append(content + '\n\\par\n')
        elif isinstance(block, Table):
            output.append('\\par\\smallskip\n' + table_tex([[cell.text for cell in row.cells] for row in block.rows]) + '\n\\par\\smallskip')
    return '\n'.join(output)
